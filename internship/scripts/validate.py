import shutil
import subprocess
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
MARKDOWN = ROOT / "internship" / "de-cuong-thuc-tap.md"
DOCX = ROOT / "internship" / "de-cuong-thuc-tap.docx"
DOC = ROOT / "internship" / "de-cuong-thuc-tap.doc"

REQUIRED = (
    "Nguyễn Thành Vinh",
    "23IT313",
    "ThS. Trần Đình Sơn",
    "tdson@vku.udn.vn",
    "1. Mục tiêu thực tập",
    "2. Đề tài",
    "3. Mô tả chi tiết",
    "4. Kế hoạch thực hiện",
    "03/08/2026",
    "09/08/2026",
)
FORBIDDEN = ("EKS", "Helm", "Argo CD", "GitOps", "thương mại điện tử", "Đinh Minh Khoa", "23IT128")


def assert_document(text, source):
    for expected in REQUIRED:
        assert expected in text, f"{source}: missing {expected!r}"
    for forbidden in FORBIDDEN:
        assert forbidden not in text, f"{source}: forbidden {forbidden!r}"
    assert text.count("2026 –") >= 8, f"{source}: expected eight schedule rows"


def main():
    assert MARKDOWN.exists(), f"Missing Markdown source: {MARKDOWN}"
    assert_document(MARKDOWN.read_text(encoding="utf-8"), MARKDOWN)

    assert DOCX.exists(), f"Missing generated .docx: {DOCX}"
    document = Document(DOCX)
    docx_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    assert_document(docx_text, DOCX)

    if DOC.exists():
        antiword = shutil.which("antiword")
        assert antiword, "antiword is required to validate the .doc export"
        result = subprocess.run([antiword, "-m", "UTF-8", str(DOC)], capture_output=True, text=True, check=True)
        assert_document(result.stdout, DOC)
        print("Internship proposal Markdown, DOCX, and DOC validated")
    else:
        print("Internship proposal Markdown and DOCX validated; DOC conversion is optional and was not produced")


if __name__ == "__main__":
    main()
